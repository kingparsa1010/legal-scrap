"""
Usage
-----
    pip install playwright beautifulsoup4 python-docx
    playwright install chromium

    python scrape_brands.py



    
ipm.ssaa.ir brand-declaration extractor
========================================
1. LOGIN (manual: a visible Chromium window opens.
   You log in and solve the captcha yourself. The script does not touch
   the login form or the captcha in any way - it just waits for you.
2. SCRAPE (automated): once you confirm you're logged in, the script reuses
   that same authenticated browser session to walk through the paginated
   declaration list, page by page, and grabs the raw HTML of each page.
3. PARSE: each page's HTML is parsed with BeautifulSoup to pull out every
   field per brand (declaration number, title, office, owner(s), date,
   status, image link, detail link).
4. STORE + EXPORT: results are kept in a local JSON file (data_store.json)
   keyed by declaration number, so re-running the script only appends
   brands you don't already have - it never duplicates entries in the
   Word file. The Word document gets one section per brand.

Because the list is sorted newest-first (confirmed by the descending dates
on page 1), once a page contains *zero* brands you don't already have, the
script assumes everything after that page is also already known and stops
early. Pass --full-rescan to disable that shortcut and always walk every
page (slower, but useful if you ever suspect the sort order changed).
"""

import argparse
import json
import re
import sys
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

try:
    from playwright.sync_api import sync_playwright
except ImportError:
    sys.exit(
        "Playwright is not installed. Run:\n"
        "    pip install playwright\n"
        "    playwright install chromium"
    )

BASE_LIST_URL = (
    "https://ipm.ssaa.ir/CA/DeclarationList"
    "?DeclarationType=&decNo=&FromDate=&ToDate=&DeclarationState="
    "&dTitle=&DeclarationSort=1&TotalRecord=1001&p={page}"
)
START_URL = "https://ipm.ssaa.ir/CA/DeclarationList"


def set_paragraph_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def set_run_rtl(run):
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rPr.append(rtl)
    run.font.name = "Tahoma"
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), "Tahoma")


def add_rtl_paragraph(doc, text, bold=False, size=11, heading_level=None):
    if heading_level:
        p = doc.add_heading(level=heading_level)
    else:
        p = doc.add_paragraph()
    set_paragraph_rtl(p)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    set_run_rtl(run)
    return p


# --------------------------------------------------------------------------
# Parsing
# --------------------------------------------------------------------------
def parse_declaration_page(html):
    """Return a list of dicts, one per brand row, for a single list page."""
    soup = BeautifulSoup(html, "html.parser")
    tbody = soup.find("tbody", id="DataList")
    if not tbody:
        return []

    rows = []
    for tr in tbody.find_all("tr", recursive=False):
        tds = tr.find_all("td", recursive=False)
        if len(tds) < 8:
            continue

        row_num = tds[0].get_text(strip=True)

        img_a = tds[1].find("a")
        image_url = img_a["href"].strip() if img_a and img_a.has_attr("href") else ""

        dec_a = tds[2].find("a")
        dec_no_display = dec_a.get_text(strip=True) if dec_a else tds[2].get_text(strip=True)
        detail_url = dec_a["href"].strip() if dec_a and dec_a.has_attr("href") else ""
        m = re.search(r"DecNo=(\d+)", detail_url)
        dec_no_id = m.group(1) if m else dec_no_display

        title = tds[3].get_text(strip=True)
        office = tds[4].get_text(strip=True)

        owner_b = tds[5].find("b")
        percent = ""
        owner_name = ""
        if owner_b:
            percent_i = owner_b.find("i")
            percent = percent_i.get_text(strip=True) if percent_i else ""
            owner_name = owner_b.get_text(separator=" ", strip=True)
            if percent:
                owner_name = owner_name.replace(percent, "").strip()
        else:
            owner_name = tds[5].get_text(strip=True)

        date_text = tds[6].get_text(separator="|", strip=True)
        date_parts = [p for p in date_text.split("|") if p]
        reg_date = date_parts[0] if len(date_parts) > 0 else ""
        reg_time = date_parts[1] if len(date_parts) > 1 else ""

        status_span = tds[7].find("span")
        status = status_span.get_text(strip=True) if status_span else tds[7].get_text(strip=True)

        rows.append(
            {
                "dec_no_id": dec_no_id,
                "row_num": row_num,
                "dec_no_display": dec_no_display,
                "title": title,
                "office": office,
                "owner_name": owner_name,
                "owner_percent": percent,
                "reg_date": reg_date,
                "reg_time": reg_time,
                "status": status,
                "image_url": image_url,
                "detail_url": detail_url,
            }
        )
    return rows


def is_logged_in(page):
    """userData div (with the account name + logout link) only appears
    when authenticated; the login form appears otherwise."""
    return page.locator("#userData").count() > 0


# --------------------------------------------------------------------------
# Word document
# --------------------------------------------------------------------------
def load_or_create_doc(path: Path):
    if path.exists():
        return Document(str(path))
    doc = Document()
    add_rtl_paragraph(doc, "لیست اظهارنامه‌های ثبت‌شده", heading_level=1)
    return doc


def append_brand_section(doc, entry):
    add_rtl_paragraph(doc, entry["title"] or "(بدون عنوان)", bold=True, size=14, heading_level=2)
    fields = [
        ("شماره اظهارنامه / کد رهگیری", entry["dec_no_display"]),
        ("اداره", entry["office"]),
        ("مالک/مالکین", f'{entry["owner_name"]}  {entry["owner_percent"]}'.strip()),
        ("تاریخ ثبت", entry["reg_date"]),
        ("ساعت ثبت", entry["reg_time"]),
        ("وضعیت", entry["status"]),
    ]
    for label, value in fields:
        p = add_rtl_paragraph(doc, f"{label}: {value}")
    if entry["image_url"]:
        add_rtl_paragraph(doc, f'تصویر اظهارنامه: {entry["image_url"]}')
    if entry["detail_url"]:
        add_rtl_paragraph(doc, f'لینک جزئیات: {entry["detail_url"]}')
    doc.add_paragraph()  # spacer


# --------------------------------------------------------------------------
# Store
# --------------------------------------------------------------------------
def load_store(path: Path):
    if path.exists():
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_store(path: Path, store: dict):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(store, f, ensure_ascii=False, indent=2)


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", default="brands.docx", help="Word file to create/update")
    parser.add_argument("--store", default="data_store.json", help="JSON cache of already-seen brands")
    parser.add_argument("--max-pages", type=int, default=500, help="hard safety cap on pages walked")
    parser.add_argument(
        "--full-rescan",
        action="store_true",
        help="Ignore the early-stop shortcut and walk every page every run",
    )
    args = parser.parse_args()

    output_path = Path(args.output)
    store_path = Path(args.store)
    store = load_store(store_path)
    is_first_run = len(store) == 0

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        page = browser.new_page()
        page.goto(START_URL)

        print("\n" + "=" * 70)
        print("یک پنجره مرورگر باز شد.")
        print("لطفاً به‌صورت دستی وارد حساب کاربری خود شوید و کپچا را حل کنید.")
        print("وقتی صفحه‌ی لیست اظهارنامه‌ها را دیدید، به اینجا برگردید و Enter بزنید.")
        print("=" * 70)
        input("Press Enter once you are logged in and the brand list is visible... ")

        while not is_logged_in(page):
            print("هنوز لاگین تشخیص داده نشد. لطفاً مطمئن شوید وارد شده‌اید، سپس دوباره Enter بزنید.")
            input("Press Enter to check again... ")

        print("Login detected. Starting automated scrape...\n")

        doc = load_or_create_doc(output_path)
        new_count = 0
        page_num = 1

        while page_num <= args.max_pages:
            url = BASE_LIST_URL.format(page=page_num)
            page.goto(url)
            html = page.content()
            rows = parse_declaration_page(html)

            if not rows:
                print(f"Page {page_num}: no rows found, stopping.")
                break

            new_on_this_page = 0
            for entry in rows:
                if entry["dec_no_id"] in store:
                    continue
                append_brand_section(doc, entry)
                store[entry["dec_no_id"]] = entry
                new_on_this_page += 1
                new_count += 1

            print(f"Page {page_num}: {len(rows)} rows, {new_on_this_page} new.")

            if not args.full_rescan and not is_first_run and new_on_this_page == 0:
                print("No new brands on this page - assuming the rest is already known. Stopping early.")
                break

            page_num += 1

        browser.close()

    if new_count > 0:
        doc.save(str(output_path))
        save_store(store_path, store)
        print(f"\nDone. {new_count} new brand(s) added to {output_path}.")
    else:
        print("\nDone. No new brands found - Word file left unchanged.")


if __name__ == "__main__":
    main()

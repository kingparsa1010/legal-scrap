"""
تولیدکننده گزارش‌های اکسل با استایل استاندارد و پشتیبانی کامل راست‌به‌چپ (RTL)
Excel Exporter — mode-aware sheets:
  Mode 1: ALL SCRAPED RECORDS + ACTIVE ONLY
  Mode 2: CAPTURED CLASS 5 + MATCHED
"""
import os
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from scanner import BASE_DIR as ROOT_DIR
DEFAULT_REPORTS_DIR = os.path.join(ROOT_DIR, "reports")


HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
HEADER_FONT = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
HIGH_FILL = PatternFill(start_color="FFD2D2", end_color="FFD2D2", fill_type="solid")
HIGH_FONT = Font(name="Tahoma", size=10, bold=True, color="9C0006")
MED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
MED_FONT = Font(name="Tahoma", size=10, bold=True, color="9C6500")
CELL_FONT = Font(name="Tahoma", size=10)
LINK_FONT = Font(name="Tahoma", size=10, color="0000FF", underline="single")
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

RECORD_HEADERS = [
    "ردیف", "عنوان علامت", "شماره اظهارنامه", "تاریخ اظهارنامه",
    "وضعیت پرونده", "ثبت قطعی", "طبقات بین‌المللی", "لینک پرونده در سامانه"
]

MATCHED_HEADERS = [
    "ردیف", "درصد شباهت", "سطح ریسک",
    "عنوان (data_store)", "شماره اظهارنامه", "تاریخ", "لینک",
    "عنوان (کلاس ۵)", "شماره اظهارنامه", "تاریخ", "لینک"
]


def _style_header_row(ws, headers: List[str]):
    for col_idx, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER


def _add_record_sheet(wb, title: str, records: List[Dict[str, Any]]):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].rightToLeft = True
    ws.append(RECORD_HEADERS)
    _style_header_row(ws, RECORD_HEADERS)

    for idx, item in enumerate(records, start=1):
        active_fa = "ثبت قطعی" if item.get('is_active') else "—"
        classes = item.get('classes', '')
        if isinstance(classes, list):
            classes = ','.join(map(str, classes))
        row_data = [
            idx,
            item.get('title', ''),
            item.get('dec_no_display') or item.get('declaration_no', ''),
            item.get('declaration_date') or item.get('reg_date', ''),
            item.get('status', ''),
            active_fa,
            classes,
            item.get('detail_url', ''),
        ]
        ws.append(row_data)
        row_num = ws.max_row
        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=c_idx)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = RIGHT if c_idx == 2 else CENTER
            if c_idx == 6 and item.get('is_active'):
                cell.font = Font(name="Tahoma", size=10, bold=True, color="15803D")
            if c_idx == 8 and item.get('detail_url'):
                cell.hyperlink = item.get('detail_url')
                cell.font = LINK_FONT
    return ws


def _add_matched_sheet(wb, title: str, matches: List[Dict[str, Any]]):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].rightToLeft = True
    ws.append(MATCHED_HEADERS)
    _style_header_row(ws, MATCHED_HEADERS)

    for idx, m in enumerate(matches, start=1):
        score = float(m.get('score', 0.0))
        risk = m.get('risk', 'medium')
        risk_fa = "خطر بالا" if risk == 'high' else "خطر متوسط"
        row_data = [
            idx,
            f"{score:.1f}%",
            risk_fa,
            m.get('title_a', ''),
            m.get('declaration_no_a', ''),
            m.get('declaration_date_a', ''),
            m.get('detail_url_a', ''),
            m.get('title_b', ''),
            m.get('declaration_no_b', ''),
            m.get('declaration_date_b', ''),
            m.get('detail_url_b', ''),
        ]
        ws.append(row_data)
        row_num = ws.max_row
        for c_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=row_num, column=c_idx)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            cell.alignment = RIGHT if c_idx in (4, 8) else CENTER
            if c_idx in (2, 3):
                if risk == 'high':
                    cell.fill = HIGH_FILL
                    cell.font = HIGH_FONT
                else:
                    cell.fill = MED_FILL
                    cell.font = MED_FONT
            if c_idx in (7, 11) and row_data[c_idx - 1]:
                cell.hyperlink = row_data[c_idx - 1]
                cell.font = LINK_FONT
    return ws


def _autofit(wb):
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 60)


def export_scan_to_excel(
    scan_info: Dict[str, Any],
    all_records: Optional[List[Dict[str, Any]]] = None,
    active: Optional[List[Dict[str, Any]]] = None,
    class5_records: Optional[List[Dict[str, Any]]] = None,
    matched: Optional[List[Dict[str, Any]]] = None,
    output_dir: Optional[str] = None,
    prefix: str = "گزارش_اسکن",
) -> str:
    """
    تولید فایل اکسل متناسب با حالت:
      حالت ۱ → برگه‌های ALL + ACTIVE
      حالت ۲ → برگه‌های CLASS5 (+ MATCHED در صورت وجود نتیجه مقایسه)
    حداقل یکی از ورودی‌ها باید مقدار داشته باشد.
    """
    if output_dir is None:
        output_dir = DEFAULT_REPORTS_DIR
    os.makedirs(output_dir, exist_ok=True)
    scan_date_str = str(scan_info.get('scan_date', 'unknown')).replace('/', '-').replace(':', '-')
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{prefix}_{scan_date_str}_{timestamp}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb = openpyxl.Workbook()
    # حذف شیت پیش‌فرض خالی
    wb.remove(wb.active)

    if all_records is not None:
        _add_record_sheet(wb, "ALL SCRAPED RECORDS", all_records)
    if active is not None:
        _add_record_sheet(wb, "ACTIVE ONLY", active)
    if class5_records is not None:
        _add_record_sheet(wb, "CAPTURED CLASS 5", class5_records)
    if matched is not None:
        _add_matched_sheet(wb, "MATCHED", matched)

    if not wb.worksheets:
        _add_record_sheet(wb, "ALL SCRAPED RECORDS", [])

    _autofit(wb)
    wb.save(filepath)
    return filepath

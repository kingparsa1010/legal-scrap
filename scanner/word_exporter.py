"""
تولیدکننده و به‌روزرسان خروجی Word جدول‌بندی‌شده راست‌به‌چپ (RTL)
Word Table Exporter for Trademark Declarations with RTL alignment, embedded images, and clickable links.
"""
import os
import io
import re
from typing import List, Dict, Any, Optional
from xml.sax.saxutils import escape as xml_escape
import requests
import urllib3

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

BASE_URL = "https://ipm.ssaa.ir"


def set_doc_rtl(doc: Document):
    """تنظیم جهت کلی سند Word به راست‌به‌چپ."""
    for section in doc.sections:
        sectPr = section._sectPr
        bidi = OxmlElement("w:bidi")
        sectPr.append(bidi)


def set_table_rtl(table):
    """تنظیم جهت و چینش جدول به راست‌به‌چپ (RTL)."""
    tblPr = table._tbl.tblPr
    bidiVisual = OxmlElement("w:bidiVisual")
    tblPr.append(bidiVisual)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def set_cell_background(cell, hex_color: str):
    """رنگ‌آمیزی پس‌زمینه خانه جدول."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    """تنظیم پدینگ خانه‌های جدول (بر حسب twips)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_rtl_run(paragraph, text: str, bold: bool = False, size: int = 9, color: Optional[RGBColor] = None):
    """افزودن یک تکست‌ران با قلم تاهما و پشتیبانی فونت فارسی RTL."""
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Tahoma"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color

    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement("w:rtl")
    rPr.append(rtl)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), "Tahoma")
    return run


def add_hyperlink(paragraph, url: str, text: str, size: int = 9, color_hex: str = "1D4ED8"):
    """افزودن لینک کلیک‌خور در پاراگراف Word با فونت و قلم مناسب."""
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    safe_text = xml_escape(text, {'"': '&quot;'})
    hyperlink = parse_xml(
        f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'
    )
    new_run = parse_xml(
        f'<w:r {nsdecls("w")}><w:rPr><w:rtl/><w:rFonts w:cs="Tahoma"/><w:color w:val="{color_hex}"/><w:sz w:val="{size*2}"/><w:u w:val="single"/></w:rPr><w:t xml:space="preserve">{safe_text}</w:t></w:r>'
    )
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _safe_float(value: Any, default: float = 0.0) -> float:
    """تبدیل امن به float؛ در صورت مقدار خراب به‌جای کرش، مقدار پیش‌فرض برمی‌گرداند."""
    try:
        if value is None or value == "":
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def _is_valid_image(data: bytes) -> bool:
    """اعتبارسنجی سبک بایت‌های تصویر قبل از تزریق به Word (جلوگیری از کرش روی دیتای خراب)."""
    if not data or len(data) < 100:
        return False
    try:
        from PIL import Image
        with Image.open(io.BytesIO(data)) as img:
            img.verify()
        return True
    except Exception:
        # اگر Pillow در دسترس نباشد یا تصویر خراب باشد، ریسک نکن
        try:
            import PIL  # noqa: F401
            return False
        except ImportError:
            # بدون Pillow نمی‌شود اعتبارسنجی کرد؛ قبولش کن تا رفتار قبلی حفظ شود
            return True


from scanner import BASE_DIR, DATA_DIR
CACHE_DIR = os.path.join(DATA_DIR, "cache", "images")


def download_image(url: str, session: Optional[requests.Session] = None, allow_network: bool = True) -> Optional[bytes]:
    """دانلود تصویر لوگو با ذخیره‌سازی محلی جهت افزایش چشمگیر سرعت خروجی Word."""
    if not url:
        return None

    # تولید نام فایل کش از روی انتهای آدرس
    clean_name = re.sub(r'[^a-zA-Z0-9_\.]', '_', url.split('/')[-1])
    if not clean_name or len(clean_name) < 4:
        clean_name = f"img_{abs(hash(url))}.jpg"
    cache_path = os.path.join(CACHE_DIR, clean_name)

    if os.path.exists(cache_path) and os.path.getsize(cache_path) > 100:
        try:
            with open(cache_path, "rb") as f:
                return f.read()
        except Exception:
            pass

    if not allow_network:
        return None

    full_url = url if url.startswith("http") else f"{BASE_URL.rstrip('/')}/{url.lstrip('/')}"
    try:
        s = session or requests
        resp = s.get(full_url, timeout=1.5, verify=False)
        if resp.status_code == 200 and len(resp.content) > 100:
            os.makedirs(CACHE_DIR, exist_ok=True)
            with open(cache_path, "wb") as f:
                f.write(resp.content)
            return resp.content
    except Exception:
        pass
    return None



def export_declarations_to_word(
    declarations: List[Dict[str, Any]],
    output_path: str = "",
    session: Optional[requests.Session] = None,
    progress_callback=None
) -> str:
    """
    تولید یا بازنویسی فایل Word به‌صورت جدول مرتب و راست‌به‌چپ:
    - شماره‌گذاری ردیف‌ها از ۱ (جدیدترین در ردیف ۱)
    - ستون‌های منظم و خوانا
    - تصاویر بندانگشتی تعبیه‌شده
    - لینک مستقیم به سامانه
    """
    if not output_path or output_path == "brands.docx":
        output_path = os.path.join(DATA_DIR, "brands.docx")
    doc = Document()
    set_doc_rtl(doc)

    # عنوان سند
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_rtl_run(p_title, "لیست اظهارنامه‌های ثبت‌شده (مرکز مالکیت معنوی)", bold=True, size=15, color=RGBColor(15, 23, 42))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_rtl_run(p_sub, f"تعداد کل اظهارنامه‌ها: {len(declarations)} مورد | ترتیب: جدیدترین در بالای جدول", size=9.5, color=RGBColor(100, 116, 139))
    doc.add_paragraph()  # فاصله

    # ستون‌های جدول (۹ ستون)
    headers = [
        ("ردیف", Inches(0.4)),
        ("تصویر", Inches(0.9)),
        ("عنوان علامت", Inches(1.3)),
        ("شماره اظهارنامه", Inches(1.1)),
        ("مالک / مالکین", Inches(1.2)),
        ("تاریخ ثبت", Inches(0.8)),
        ("ساعت", Inches(0.5)),
        ("وضعیت", Inches(1.0)),
        ("پرونده", Inches(0.7)),
    ]

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_rtl(table)

    # سطر هدر
    hdr_cells = table.rows[0].cells
    for i, (title, width) in enumerate(headers):
        hdr_cells[i].width = width
        set_cell_background(hdr_cells[i], "1B365D")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=80, right=80)
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rtl_run(p, title, bold=True, size=9.5, color=RGBColor(255, 255, 255))

    total = len(declarations)
    network_images_allowed = True
    consecutive_image_fails = 0

    for idx, item in enumerate(declarations):
        row_cells = table.add_row().cells

        # رنگ‌بندی یک‌درمیان (Zebra striping)
        bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        # اگر در ارزیابی شباهت بالا یا متوسط داشت، رنگ شاخص بگیرد
        sim_score = _safe_float(item.get("similarity_score"), 0.0)
        risk_level = str(item.get("risk_level") or "").lower()
        if "high" in risk_level or sim_score >= 85.0:
            bg_color = "FEE2E2"  # قرمز بسیار ملایم
        elif "medium" in risk_level or sim_score >= 70.0:
            bg_color = "FEF3C7"  # کهربایی ملایم

        for c in row_cells:
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=80, bottom=80, left=60, right=60)
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # ۱. ردیف (از شماره ۱ شروع می‌شود)
        p_num = row_cells[0].paragraphs[0]
        p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rtl_run(p_num, str(idx + 1), bold=True, size=9)

        # ۲. تصویر علامت
        p_img = row_cells[1].paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_url = item.get("image_url", "")
        img_bytes = None
        if img_url:
            img_bytes = download_image(img_url, session=session, allow_network=network_images_allowed)
            if img_bytes is None and network_images_allowed:
                consecutive_image_fails += 1
                if consecutive_image_fails >= 3:
                    network_images_allowed = False
            elif img_bytes is not None:
                consecutive_image_fails = 0

        if img_bytes:
            if not _is_valid_image(img_bytes):
                img_bytes = None
        if img_bytes:
            try:
                img_stream = io.BytesIO(img_bytes)
                p_img.add_run().add_picture(img_stream, width=Inches(0.75))
            except Exception:
                add_rtl_run(p_img, "—", size=8, color=RGBColor(148, 163, 184))
        else:
            add_rtl_run(p_img, "—", size=8, color=RGBColor(148, 163, 184))

        # ۳. عنوان علامت
        p_title = row_cells[2].paragraphs[0]
        p_title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        title_text = item.get("title", "") or "(بدون عنوان)"
        add_rtl_run(p_title, title_text, bold=True, size=9)
        if sim_score >= 70.0 and item.get("matched_brand"):
            p_sub_sim = p_title.add_run(f"\n[شباهت {sim_score}% با {item.get('matched_brand')}]")
            p_sub_sim.font.name = "Tahoma"
            p_sub_sim.font.size = Pt(7.5)
            p_sub_sim.bold = True
            p_sub_sim.font.color.rgb = RGBColor(185, 28, 28) if sim_score >= 85 else RGBColor(180, 83, 9)

        # ۴. شماره اظهارنامه / رهگیری
        p_dec = row_cells[3].paragraphs[0]
        p_dec.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dec_no = item.get("dec_no_display") or item.get("declaration_no") or item.get("dec_no_id") or ""
        add_rtl_run(p_dec, str(dec_no), size=8.5)

        # ۵. مالک / مالکین
        p_owner = row_cells[4].paragraphs[0]
        p_owner.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        owner_str = str(item.get("owner_name", "")).strip()
        percent = str(item.get("owner_percent", "")).strip()
        full_owner = f"{owner_str} ({percent}%)" if percent else owner_str
        add_rtl_run(p_owner, full_owner or "—", size=8.5)

        # ۶. تاریخ ثبت
        p_date = row_cells[5].paragraphs[0]
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rtl_run(p_date, str(item.get("reg_date") or item.get("declaration_date") or "—"), size=8)

        # ۷. ساعت ثبت
        p_time = row_cells[6].paragraphs[0]
        p_time.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_rtl_run(p_time, str(item.get("reg_time") or "—"), size=8)

        # ۸. وضعیت
        p_stat = row_cells[7].paragraphs[0]
        p_stat.alignment = WD_ALIGN_PARAGRAPH.CENTER
        stat_text = str(item.get("status") or "—")
        stat_color = RGBColor(185, 28, 28) if "رد" in stat_text else RGBColor(30, 41, 59)
        add_rtl_run(p_stat, stat_text, size=8, color=stat_color)

        # ۹. لینک سامانه
        p_link = row_cells[8].paragraphs[0]
        p_link.alignment = WD_ALIGN_PARAGRAPH.CENTER
        detail_rel = item.get("detail_url") or ""
        if detail_rel:
            full_detail = detail_rel if detail_rel.startswith("http") else f"{BASE_URL.rstrip('/')}/{detail_rel.lstrip('/')}"
            add_hyperlink(p_link, full_detail, "مشاهده", size=8.5, color_hex="2563EB")
        else:
            add_rtl_run(p_link, "—", size=8)

        if progress_callback and idx % 10 == 0:
            progress_callback(idx, total)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    try:
        doc.save(output_path)
    except PermissionError as exc:
        raise PermissionError(
            f"فایل Word باز است یا دسترسی نوشتن ندارید: {output_path}\n"
            "لطفاً فایل brands.docx را در Word ببندید و دوباره تلاش کنید."
        ) from exc
    return output_path
